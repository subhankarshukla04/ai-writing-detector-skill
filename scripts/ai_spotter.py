"""ai_spotter: an explainable AI-writing check that runs locally. It reports evidence, not a verdict.

  python ai_spotter.py FILE [--human BASELINE ...] [--lm] [--binoculars] [--json]
  python ai_spotter.py --calibrate HUMAN_DIR AI_DIR [--binoculars]

FILE: .txt .md .docx .pptx (slides and speaker notes are scored separately).
--human     text known to be human (ideally the same author or genre) to compare against. Strongly advised.
--lm        GPT-2 perplexity and burstiness (GPTZero's original signals). Downloads ~500MB once.
--binoculars  Binoculars score (Hans et al., ICML 2024) with a small same-tokenizer model pair.
            Uncalibrated on small models: run --calibrate on your own human/AI samples first.

Layers, and where each rule comes from:
  1 tells        Wikipedia "Signs of AI writing" (WikiProject AI Cleanup); Kobak et al. 2024
  2 stylometry   Reinhart et al. PNAS 2025 (participial clauses, nominalisations); sentence-length variance
  3 lm           perplexity + burstiness (GPTZero's original method); Binoculars (Hans et al. 2024)
No detector is reliable enough to accuse a person. See reference/limits.md.
"""
import json
import math
import re
import sys
from pathlib import Path

ERA_VOCAB = {  # [W] AI vocabulary by period. Ordinary words on Wikipedia's newest list ("matters", "steady",
    # "challenge") are left out: alone they flag ordinary human prose.
    "2023-24": "additionally boasts bolstered delve delves delving garner intricate intricacies interplay meticulous "
               "meticulously pivotal tapestry testament vibrant",
    "2024-26": "crucial enduring emphasizing enhance enhances fostering highlighting showcasing underscore underscores",
    "2026+": "foster leverage multifaceted prioritize",
    "general": "realm robust seamless seamlessly holistic synergy unlock elevate empower game-changer cutting-edge "
               "paradigm harness nuanced myriad plethora embark furthermore moreover notably transformative "
               "groundbreaking streamline spearhead delved leveraged fostered harnessed empowered elevated unlocked "
               "underscored showcased embarked",
}
VOCAB_MIN = 2  # a unit's vocab words only count when at least this many distinct ones appear together
TELLS = {  # name: (regex, weight, source)
    "negative parallelism": (r"\bnot (?:just|only|merely|simply)\b[^.;:\n]{1,80}?\bbut\b|"
                             r"\b(?:isn't|is not|wasn't|aren't|it's not)\b[^.;:\n]{1,60}?[,;—-]\s*(?:it's|it is|they're|it was)\b", 2, "W"),
    "copula avoidance": (r"\b(?:serves as|stands as|functions as|operates as|acts as|boasts)\b", 1.5, "W"),
    "-ing analysis tail": (r",\s(?:highlighting|underscoring|emphasizing|emphasising|ensuring|reflecting|symbolizing|"
                           r"contributing to|cultivating|fostering|encompassing|showcasing)\b", 2, "W, R"),
    "undue significance": (r"\b(?:pivotal (?:moment|role)|crucial role|vital role|key role|enduring legacy|lasting impact|"
                           r"indelible|testament to|broader trends?|focal point|setting the stage|evolving landscape)\b", 2, "W"),
    "promotional": (r"\b(?:nestled|breathtaking|in the heart of|renowned|diverse array|rich (?:history|heritage|tradition)|"
                    r"commitment to excellence|vibrant community)\b", 1.5, "W"),
    "vague attribution": (r"\b(?:experts (?:say|argue|note|agree)|observers (?:say|note)|industry reports|"
                          r"some critics|many believe|studies (?:show|suggest))\b", 1.5, "W"),
    "challenges/outlook outro": (r"\b(?:despite (?:these|its|the|their) challenges|future outlook|looking ahead|"
                                 r"the road ahead|future prospects)\b", 1.5, "W"),
    "template reveal": (r"\b(?:The result|The answer|The catch|The twist|The kicker|The best part)\?|"
                        r"\b(?:Here's|Here is) (?:the thing|the kicker|the catch|what|why|how)\b[^.\n]{0,40}:", 1.5, "W"),
    "collaborative chat": (r"\b(?:i hope this helps|let me know if|feel free to|great question|"
                           r"as an ai|happy to help)\b|\bcertainly!", 3, "W"),
    "chatbot artifact": (r"oaicite|contentReference|turn0search|\[cite:\s?\d|grok_card|attributableIndex|"
                         r"utm_source=chatgpt|【\d+†", 5, "W"),
    "bold inline header": (r"^\s*(?:[-*•]|\d+\.)\s*\*\*[^*\n]{1,40}?:?\*\*:?\s", 1, "W"),
    "emoji formatting": ("[\U0001F300-\U0001FAFF✅❌✨]", 1, "W"),
    "em dash": (r"—", 0.5, "W historical"),
    "hedge stack": (r"\b(?:may|might|could)\s+(?:potentially|possibly|perhaps)\b", 1, "practice"),
}
NOMINAL = re.compile(r"\b\w{4,}(?:tion|tions|ment|ments|ness|ity|ities|ance|ence)\b", re.I)
PARTICIPLE_CLAUSE = re.compile(r"(?:^|[,;]\s)(?:\w+ing)\b", re.I | re.M)


# ---------- reading ----------
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)
        else:
            yield sh


def normalise(t):
    return t.replace("\u2019", "'").replace("\u2018", "'").replace("\u201c", '"').replace("\u201d", '"')


def extract(path):
    return [(u, k, normalise(t)) for u, k, t in _extract(path)]


def _extract(path):
    p = str(path)
    low = p.lower()
    if not low.endswith((".pptx", ".docx", ".txt", ".md")):
        sys.exit(f"unsupported file type: {p} (use .txt, .md, .docx or .pptx)")
    if low.endswith(".pptx"):
        from pptx import Presentation
        units = []
        for i, s in enumerate(Presentation(p).slides, 1):
            parts = []
            for sh in walk(s.shapes):
                if sh.has_text_frame and sh.text_frame.text.strip():
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    parts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
            units.append((f"slide {i}", "slide", "\n".join(parts)))
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                units.append((f"slide {i} notes", "prose", s.notes_slide.notes_text_frame.text))
        return units
    if low.endswith(".docx"):
        import docx
        doc = docx.Document(p)
        text = "\n\n".join([x.text for x in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    else:
        text = Path(p).read_text(encoding="utf-8", errors="ignore")
    return [(f"para {i}", "prose", t) for i, t in enumerate(re.split(r"\n\s*\n", text), 1) if t.strip()]


def sentences(t):
    flat = re.sub(r"\s+", " ", re.sub(r"(?m)^\s*(?:[#>*•-]+|\d+\.)\s*", "", t))
    return [s for s in re.split(r"(?<=[.!?])\s+", flat) if len(s.split()) >= 4]


def words(t):
    return re.findall(r"[a-z][a-z'-]*", t.lower())


# ---------- layer 1: tells ----------
def tells(text):
    vocab = {w: era for era, ws in ERA_VOCAB.items() for w in ws.split()}
    hits = [w for w in words(text) if w in vocab]
    found = [(f"vocab '{w}' ({vocab[w]})", 1.0, "W") for w in hits] if len(set(hits)) >= VOCAB_MIN else []
    for name, (rx, wt, src) in TELLS.items():
        for m in re.finditer(rx, text, re.I | re.M):
            found.append((f"{name}: '{m.group(0)[:50].strip()}'", wt, src))
    for t in re.findall(r"\b[\w-]+(?: [\w-]+)?, [\w-]+(?: [\w-]+)?,? and [\w-]+\b", text):
        found.append((f"rule of three: '{t[:50]}'", 0.5, "W"))
    return found


# ---------- layer 2: stylometry ----------
def stylometry(text):
    ss, ws = sentences(text), words(text)
    out = {"words": len(ws), "sentences": len(ss)}
    lens = [len(s.split()) for s in ss]
    if len(lens) >= 3:
        m = sum(lens) / len(lens)
        out["sent_len_mean"] = m
        out["sent_len_cv"] = (sum((x - m) ** 2 for x in lens) / len(lens)) ** 0.5 / m
    if len(ws) > 60:
        win = 50
        out["mattr"] = sum(len(set(ws[i:i + win])) / win for i in range(len(ws) - win)) / (len(ws) - win)
        out["nominal_per_100w"] = 100 * len(NOMINAL.findall(text)) / len(ws)
        out["participle_clauses_per_100w"] = 100 * len(PARTICIPLE_CLAUSE.findall(text)) / len(ws)
    return out


# ---------- layer 3: language models ----------
def gpt2():
    import torch
    from transformers import GPT2LMHeadModel, GPT2TokenizerFast
    tok = GPT2TokenizerFast.from_pretrained("gpt2")
    model = GPT2LMHeadModel.from_pretrained("gpt2").eval()

    def ppl(sent):
        ids = tok(sent, return_tensors="pt").input_ids[:, :512]
        if ids.shape[1] < 4:
            return None
        with torch.no_grad():
            return math.exp(model(ids, labels=ids).loss.item())
    return ppl


def lm_stats(text, ppl):
    ps = [(p, s) for s in sentences(text)[:400] if (p := ppl(s))]
    if len(ps) < 3:
        return {}, []
    lp = [math.log(p) for p, _ in ps]
    mu = sum(lp) / len(lp)
    vals = sorted(p for p, _ in ps)
    return ({"ppl_median": vals[len(vals) // 2],
             "ppl_burstiness": (sum((x - mu) ** 2 for x in lp) / len(lp)) ** 0.5}, sorted(ps)[:8])


OBSERVER, PERFORMER = "Qwen/Qwen2.5-0.5B", "Qwen/Qwen2.5-0.5B-Instruct"


def binoculars():
    """Binoculars score = log-perplexity under the performer / cross-perplexity (observer probs vs performer logits).
    Same formula as github.com/ahans30/Binoculars, with a small pair that shares a tokenizer. Lower = more machine-like."""
    import torch
    from transformers import AutoModelForCausalLM, AutoTokenizer
    tok = AutoTokenizer.from_pretrained(OBSERVER)
    obs = AutoModelForCausalLM.from_pretrained(OBSERVER, dtype=torch.float32).eval()
    per = AutoModelForCausalLM.from_pretrained(PERFORMER, dtype=torch.float32).eval()

    def score(text):
        ids = tok(text, return_tensors="pt", truncation=True, max_length=512).input_ids
        if ids.shape[1] < 16:
            return None
        with torch.no_grad():
            lo, lp_ = obs(ids).logits[0, :-1], per(ids).logits[0, :-1]
        target = ids[0, 1:]
        log_ppl = torch.nn.functional.cross_entropy(lp_, target).item()
        x_ppl = (-(torch.softmax(lo, -1) * torch.log_softmax(lp_, -1)).sum(-1)).mean().item()
        return log_ppl / x_ppl
    return score


# ---------- report ----------
def lean(value, human, higher_is_human=True, tol=0.10):
    if value is None or human is None:
        return "no baseline"
    if abs(value - human) <= tol * abs(human):
        return "close to the human baseline"
    lower = value < human
    if higher_is_human:
        return "lower than baseline (AI-leaning in our test)" if lower else "higher than baseline (human-leaning)"
    return "higher than baseline" if not lower else "lower than baseline"


CAVEAT = ("Evidence for an editor, not proof of authorship. Tells say what to rewrite. Model signals only mean "
          "something against a human baseline from the same genre. Detectors misfire on non-native and very formal "
          "writing (Liang et al. 2023). Never use this as the sole basis for a decision about a person.")


def analyse(path, humans, use_lm, use_bino, as_json):
    units = extract(path)
    prose = "\n".join(t for _, k, t in units if k == "prose") or "\n".join(t for _, _, t in units)
    allt = "\n".join(t for _, _, t in units)
    res = {"file": str(path), "stylometry": stylometry(prose), "baselines": {}}
    base_texts = {(Path(h).name if [Path(x).name for x in humans].count(Path(h).name) == 1 else str(h)): "\n".join(t for _, _, t in extract(h))[:60000] for h in humans}
    for name, bt in base_texts.items():
        res["baselines"][name] = stylometry(bt)
        res["baselines"][name]["tells_per_1k"] = 1000 * sum(f[1] for f in tells(bt)) / max(1, len(words(bt)))

    found = [(u, *f) for u, k, t in units for f in tells(t)]
    weight = sum(f[2] for f in found)
    res["tells"] = {"count": len(found), "weighted_per_1k_words": 1000 * weight / max(1, len(words(allt)))}

    worst = []
    if use_lm:
        ppl = gpt2()
        st, worst = lm_stats(prose, ppl)
        res["stylometry"].update(st)
        for name, bt in base_texts.items():
            res["baselines"][name].update(lm_stats(bt, ppl)[0])
    if use_bino:
        b = binoculars()
        res["binoculars"] = b(prose[:6000])
        for name, bt in base_texts.items():
            res["baselines"][name]["binoculars"] = b(bt[:6000])

    res["caveat"] = CAVEAT
    if res["stylometry"]["words"] < 300:
        res["warning"] = f"only {res['stylometry']['words']} words: statistics are unreliable below about 300"
    if as_json:
        res["tell_list"] = [{"where": u, "tell": w, "weight": wt, "source": s} for u, w, wt, s in found]
        print(json.dumps(res, indent=2, default=float))
        return

    st = res["stylometry"]
    print(f"# AI-writing evidence report: {Path(path).name}\n")
    print(f"{st['words']} words, {st['sentences']} sentences scored.\n")
    print(f"> {CAVEAT}\n")
    if "warning" in res:
        print(f"**Warning: {res['warning']}.**\n")
    print("## 1. Documented tells (Wikipedia Signs of AI writing, Kobak et al.)\n")
    d = res["tells"]["weighted_per_1k_words"]
    bd = {b: v["tells_per_1k"] for b, v in res["baselines"].items()}
    cmp_ = ("human baseline " + "; ".join(f"{b}: {v:.1f}" for b, v in bd.items())) if bd else "no human baseline given"
    print(f"{len(found)} tells, weighted density {d:.1f} per 1,000 words ({cmp_}).\n")
    cur = None
    for u, w, wt, s in found:
        if u != cur:
            print(f"**{u}**")
            cur = u
        print(f"- {w}  _[{s}]_")

    print("\n## 2. Style and model signals\n")
    # (key, higher_is_human or None, label). None = direction did not hold up in our calibration
    # (README: 15 pre-2022 human passages vs 10 passages from a 2026 frontier model), so no lean is shown.
    keys = [("sent_len_cv", True, "sentence-length variation (separated 21/25 in our small test)"),
            ("binoculars", True, "Binoculars score, Qwen2.5-0.5B pair (separated 19/25 in our small test)"),
            ("mattr", None, "lexical diversity (MATTR-50)"),
            ("participle_clauses_per_100w", None, "participial clauses / 100 words (Reinhart 2025: LLMs 2-5x)"),
            ("nominal_per_100w", None, "nominalisations / 100 words (Reinhart 2025: LLMs 1.5-2x)"),
            ("ppl_median", None, "GPT-2 perplexity, median sentence"),
            ("ppl_burstiness", None, "perplexity burstiness")]
    bnames = list(res["baselines"])
    cols = ["signal", "this text", *bnames, "reading"]
    print("| " + " | ".join(cols) + " |")
    print("|" + "---|" * len(cols))
    for k, hi, label in keys:
        v = st.get(k, res.get(k) if k == "binoculars" else None)
        if v is None:
            continue
        bv = [res["baselines"][b].get(k) for b in bnames]
        ref = next((x for x in bv if x is not None), None)
        cells = [label, f"{v:.2f}", *(f"{x:.2f}" if x is not None else "-" for x in bv),
                 lean(v, ref, hi) if hi is not None else "context only"]
        print("| " + " | ".join(cells) + " |")
    if not bnames:
        print("\nNo --human baseline given, so the signals above can't be read as leaning either way. "
              "Add a sample of known-human text in the same genre.")
    if worst:
        print("\n## 3. Most predictable sentences (check these by hand)\n")
        for p, s in worst:
            print(f"- ppl {p:.0f}: {s[:150]}")
    print("\nFor a judgement, add the blind-judge step (reference/judge-prompt.md).")


def calibrate(hdir, adir, use_bino):
    files = lambda d: [p for p in sorted(Path(d).rglob("*")) if p.suffix.lower() in (".txt", ".md", ".docx", ".pptx")]
    ppl = gpt2()
    b = binoculars() if use_bino else None
    rows = []
    for label, d in (("human", hdir), ("ai", adir)):
        for f in files(d):
            t = "\n".join(x for _, _, x in extract(f))
            st = stylometry(t)
            st.update(lm_stats(t, ppl)[0])
            if b:
                st["binoculars"] = b(t[:6000])
            st["tells_per_1k"] = 1000 * sum(x[1] for x in tells(t)) / max(1, len(words(t)))
            rows.append((label, f.name, st))
    keys = ["tells_per_1k", "sent_len_cv", "ppl_median", "ppl_burstiness"] + (["binoculars"] if b else [])
    print("| file | label | " + " | ".join(keys) + " |\n|---|---|" + "---|" * len(keys))
    for label, name, st in rows:
        print(f"| {name} | {label} | " + " | ".join(f"{st.get(k, float('nan')):.2f}" for k in keys) + " |")
    print()
    for k in keys:
        h = sorted(st[k] for l, _, st in rows if l == "human" and st.get(k) is not None)
        a = sorted(st[k] for l, _, st in rows if l == "ai" and st.get(k) is not None)
        if not h or not a:
            continue
        cands = []
        for t in h + a:
            cands.append((sum(x > t for x in h) + sum(x <= t for x in a), t, "lower = AI"))
            cands.append((sum(x <= t for x in h) + sum(x > t for x in a), t, "higher = AI"))
        acc, t, direction = max(cands)
        print(f"{k}: human median {h[len(h)//2]:.2f}, AI median {a[len(a)//2]:.2f}; best split {t:.2f} "
              f"({direction}) classifies {acc}/{len(h)+len(a)} on this set")
    print("\nUse these splits only for this genre and these models. Small samples overfit.")


if __name__ == "__main__":
    a = sys.argv
    if len(a) < 2 or a[1] in ("-h", "--help"):
        sys.exit(__doc__)
    if "--calibrate" in a:
        i = a.index("--calibrate")
        calibrate(a[i + 1], a[i + 2], "--binoculars" in a)
    else:
        humans = [a[i + 1] for i, x in enumerate(a) if x == "--human"]
        analyse(a[1], humans, "--lm" in a, "--binoculars" in a, "--json" in a)
